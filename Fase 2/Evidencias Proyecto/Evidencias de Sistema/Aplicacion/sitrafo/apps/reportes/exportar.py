"""Exportacion de reportes a Excel y PDF (RF-REP-06)."""
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import Paragraph

from apps.common.pdf import ESTILOS, construir, espacio, numero, tabla


def _texto(valor) -> str:
    if isinstance(valor, Decimal):
        return numero(valor, 2)
    return "" if valor is None else str(valor)


def a_excel(reporte: dict, periodo: str) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = reporte["titulo"][:31]
    hoja.append([reporte["titulo"]])
    hoja["A1"].font = Font(bold=True, size=14)
    hoja.append([f"Periodo: {periodo}"])
    hoja.append([])
    hoja.append(reporte["columnas"])
    fila_encabezado = hoja.max_row
    for celda in hoja[fila_encabezado]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1B242B")
    for fila in reporte["filas"]:
        hoja.append([float(v) if isinstance(v, Decimal) else v for v in fila])
    if reporte.get("totales"):
        hoja.append([float(v) if isinstance(v, Decimal) else v for v in reporte["totales"]])
        for celda in hoja[hoja.max_row]:
            celda.font = Font(bold=True)
    for i, _ in enumerate(reporte["columnas"], start=1):
        hoja.column_dimensions[get_column_letter(i)].width = 20
    memoria = BytesIO()
    libro.save(memoria)
    return memoria.getvalue()


def a_pdf(reporte: dict, periodo: str) -> bytes:
    columnas = len(reporte["columnas"])
    ancho = 245 / columnas
    filas = [reporte["columnas"]] + [[_texto(v) for v in f] for f in reporte["filas"]]
    if reporte.get("totales"):
        filas.append([f"<b>{_texto(v)}</b>" for v in reporte["totales"]])
    contenido = [
        Paragraph(reporte["titulo"], ESTILOS["titulo"]),
        Paragraph(f"Periodo: {periodo}. {reporte['descripcion']}", ESTILOS["subtitulo"]),
        tabla(filas, [ancho] * columnas) if len(filas) > 1 else
        Paragraph("Sin datos en el periodo.", ESTILOS["nota"]),
        espacio(),
    ]
    return construir(reporte["titulo"], contenido, pagina=landscape(A4))
